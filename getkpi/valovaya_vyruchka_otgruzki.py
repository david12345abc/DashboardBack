"""
Факт отгрузок KD-M2 = «Итого выручка» отчёта 1С «Валовая прибыль предприятия».

Живой источник — OData AccumulationRegister_ВыручкаИСебестоимостьПродаж_RecordType
(СуммаВыручки), те же отборы, что в выгрузке отчёта:
  • коммерческие отделы + ликвидированные дилерские;
  • без комиссионного товара и партнёра «Наше предприятие»;
  • контрагенты не из списка отчёта (АВИОН, АЛМАЗ ООО Рабочий,
    Турбулентность-Дон ООО, ТУРБУЛЕНТНОСТЬ-ДОН ООО НПО).

Если в getkpi/temp лежит файл «валовая {месяц}.xlsx», блок
«По подразделениям» / колонка «Выручка» / строка «Итого»
читаются по подписям ячеек (без шаблона колонок) и подменяют месяц.
"""
from __future__ import annotations

import json
import logging
import re
from datetime import date, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests
from requests.auth import HTTPBasicAuth

from comdir.common import MONTH_RU, name_to_odata
from getkpi.commercial_department_aliases import (
    COMMERCIAL_DEPT_ALIASES,
    DEALER_SALES_DEPT,
    normalize_commercial_dept_guid,
)
from getkpi.commercial_tiles import DEPT_GUID_TO_DZ_NAME
from getkpi.odata_http import request_with_retry

logger = logging.getLogger(__name__)

BASE = "http://192.168.2.229:81/erp_pm/odata/standard.odata"
AUTH = HTTPBasicAuth("odata.user", "npo852456")

COMMERCIAL_GUIDS: frozenset[str] = frozenset(DEPT_GUID_TO_DZ_NAME)
DEPT_FILTER_GUIDS: frozenset[str] = COMMERCIAL_GUIDS | frozenset(
    guid
    for guid, dest in COMMERCIAL_DEPT_ALIASES.items()
    if dest == DEALER_SALES_DEPT
)

NASHE_PARTNER = "5fc811cd-6251-11e7-812d-001e67112509"
EMPTY = "00000000-0000-0000-0000-000000000000"
REG_VP = "AccumulationRegister_ВыручкаИСебестоимостьПродаж_RecordType"
ANALYTICS = "Catalog_КлючиАналитикиУчетаПоПартнерам"
CONTRACTORS = "Catalog_Контрагенты"
PAGE = 5000
ANALYTICS_BATCH = 20
CACHE_VERSION = 1

EXCLUDE_CONTRACTOR_EXACT = frozenset({
    "АВИОН ООО",
    "АЛМАЗ ООО Рабочий",
    "Турбулентность-Дон ООО",
})
EXCLUDE_CONTRACTOR_PREFIX = (
    "АВИОН",
    "АЛМАЗ ООО Рабочий",
    "ТУРБУЛЕНТНОСТЬ-ДОН ООО НПО",
    "Турбулентность-ДОН ООО НПО",
)

_PARENT_DEPT_RE = re.compile(
    r"председатель|операционн|коммерческ(?:ий|ого)\s+директор",
    re.IGNORECASE,
)
_PERIOD_RE = re.compile(
    r"(\d{2})\.(\d{2})\.(\d{4})\s*[-–]\s*(\d{2})\.(\d{2})\.(\d{4})",
)


def _temp_dir() -> Path:
    return Path(__file__).resolve().parent / "temp"


def _dashboard_dir() -> Path:
    path = Path(__file__).resolve().parent / "dashboard"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _odata_cache_path(year: int, ref_month: int) -> Path:
    return _dashboard_dir() / f"vyruchka_otgruzki_odata_{year}_{ref_month:02d}.json"


def _norm_label(value: Any) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split()).strip()


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _as_amount(value: Any) -> float | None:
    if _is_number(value):
        return float(value)
    text = _norm_label(value).replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _session() -> requests.Session:
    session = requests.Session()
    session.auth = AUTH
    return session


def _odata_pages(session: requests.Session, url_base: str, *, label: str) -> list[dict]:
    rows: list[dict] = []
    skip = 0
    while True:
        url = f"{url_base}&$top={PAGE}&$skip={skip}"
        resp = request_with_retry(session, url, timeout=120, retries=4, label=label)
        if resp is None or not resp.ok:
            raise RuntimeError(f"OData {label}: HTTP {getattr(resp, 'status_code', 'drop')}")
        batch = resp.json().get("value") or []
        rows.extend(batch)
        if len(batch) < PAGE:
            break
        skip += PAGE
    return rows


def _period_window(year: int, ref_month: int) -> tuple[str, str]:
    today = date.today()
    start = date(year, 1, 1)
    if ref_month == 12:
        end = date(year + 1, 1, 1)
    else:
        end = date(year, ref_month + 1, 1)
    if year == today.year and ref_month == today.month:
        end = today + timedelta(days=1)
    return f"{start.isoformat()}T00:00:00", f"{end.isoformat()}T00:00:00"


def _dept_filter_odata() -> str:
    parts = " or ".join(f"Подразделение_Key eq guid'{guid}'" for guid in sorted(DEPT_FILTER_GUIDS))
    return f"({parts})"


def _exclude_contractor_name(name: str) -> bool:
    text = _norm_label(name)
    if not text:
        return False
    if text in EXCLUDE_CONTRACTOR_EXACT:
        return True
    return any(text.startswith(prefix) for prefix in EXCLUDE_CONTRACTOR_PREFIX)


def _fetch_contractor_names(session: requests.Session, keys: set[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    todo = [k for k in keys if k and k != EMPTY]
    for i in range(0, len(todo), ANALYTICS_BATCH):
        batch = todo[i : i + ANALYTICS_BATCH]
        flt = quote(" or ".join(f"Ref_Key eq guid'{k}'" for k in batch), safe="")
        url = (
            f"{BASE}/{CONTRACTORS}?$format=json&$top={ANALYTICS_BATCH}"
            f"&$filter={flt}&$select=Ref_Key,Description"
        )
        resp = request_with_retry(session, url, timeout=60, retries=4, label="Контрагенты")
        if resp is None or not resp.ok:
            raise RuntimeError(f"OData контрагенты: HTTP {getattr(resp, 'status_code', 'drop')}")
        for item in resp.json().get("value") or []:
            out[str(item.get("Ref_Key") or "").lower()] = _norm_label(item.get("Description"))
    return out


def _fetch_register_rows(session: requests.Session, year: int, ref_month: int) -> list[dict]:
    d_from, d_to = _period_window(year, ref_month)
    flt = quote(
        f"Period ge datetime'{d_from}' and Period lt datetime'{d_to}' "
        f"and Active eq true and {_dept_filter_odata()}",
        safe="",
    )
    sel = "Period,Active,Подразделение_Key,АналитикаУчетаПоПартнерам_Key,ТипЗапасов,СуммаВыручки"
    url = f"{BASE}/{REG_VP}?$format=json&$filter={flt}&$select={sel}"
    return _odata_pages(session, url, label="ВыручкаИСебестоимостьПродаж")


def _fetch_analytics(
    session: requests.Session,
    keys: set[str],
) -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {}
    todo = [k for k in keys if k and k != EMPTY]
    for i in range(0, len(todo), ANALYTICS_BATCH):
        batch = todo[i : i + ANALYTICS_BATCH]
        flt = quote(" or ".join(f"Ref_Key eq guid'{k}'" for k in batch), safe="")
        url = (
            f"{BASE}/{ANALYTICS}?$format=json&$top={ANALYTICS_BATCH}"
            f"&$filter={flt}&$select=Ref_Key,Партнер_Key,Контрагент,Контрагент_Type"
        )
        resp = request_with_retry(session, url, timeout=60, retries=4, label="АналитикаПартнеров")
        if resp is None or not resp.ok:
            raise RuntimeError(f"OData аналитика: HTTP {getattr(resp, 'status_code', 'drop')}")
        for item in resp.json().get("value") or []:
            ref = str(item.get("Ref_Key") or "").lower()
            out[ref] = {
                "partner": str(item.get("Партнер_Key") or "").lower(),
                "contractor": str(item.get("Контрагент") or "").lower(),
                "contractor_type": str(item.get("Контрагент_Type") or ""),
            }
    return out


def _empty_by_guid() -> dict[str, float]:
    return {guid: 0.0 for guid in COMMERCIAL_GUIDS}


def _row_month(period: str) -> int | None:
    if not period or len(period) < 7:
        return None
    try:
        return int(period[5:7])
    except ValueError:
        return None


def _fetch_odata_by_month(year: int, ref_month: int) -> dict[int, dict[str, float]]:
    session = _session()
    rows = _fetch_register_rows(session, year, ref_month)
    analytics_keys = {
        str(row.get("АналитикаУчетаПоПартнерам_Key") or "").lower()
        for row in rows
        if row.get("АналитикаУчетаПоПартнерам_Key")
    }
    analytics = _fetch_analytics(session, analytics_keys)
    contractor_keys = {
        meta.get("contractor") or ""
        for meta in analytics.values()
        if meta.get("contractor") and meta.get("contractor") != EMPTY
    }
    contractor_names = _fetch_contractor_names(session, contractor_keys)
    excluded_contractors = {
        key for key, name in contractor_names.items() if _exclude_contractor_name(name)
    }

    result = {month: _empty_by_guid() for month in range(1, int(ref_month) + 1)}
    for row in rows:
        if row.get("ТипЗапасов") == "КомиссионныйТовар":
            continue
        month = _row_month(str(row.get("Period") or ""))
        if month is None or month < 1 or month > int(ref_month):
            continue
        meta = analytics.get(str(row.get("АналитикаУчетаПоПартнерам_Key") or "").lower()) or {}
        if meta.get("partner") == NASHE_PARTNER:
            continue
        contractor = meta.get("contractor") or ""
        if contractor and contractor != EMPTY and contractor in excluded_contractors:
            continue
        guid = normalize_commercial_dept_guid(str(row.get("Подразделение_Key") or "").lower())
        if guid not in COMMERCIAL_GUIDS:
            continue
        result[month][guid] = result[month].get(guid, 0.0) + float(row.get("СуммаВыручки") or 0)
    for month, by_guid in result.items():
        result[month] = {guid: round(float(amt or 0), 2) for guid, amt in by_guid.items()}
    logger.info(
        "valovaya_vyruchka_otgruzki: OData %s-%02d строк %s итого %s",
        year, ref_month, len(rows),
        round(sum(sum(m.values()) for m in result.values()), 2),
    )
    return result


def _load_odata_cache(year: int, ref_month: int) -> dict[int, dict[str, float]] | None:
    path = _odata_cache_path(year, ref_month)
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if data.get("cache_date") != date.today().isoformat():
        return None
    if data.get("cache_version") != CACHE_VERSION:
        return None
    raw = data.get("by_month") or {}
    out: dict[int, dict[str, float]] = {}
    for key, bucket in raw.items():
        month = int(key)
        out[month] = {str(g).lower(): float(v or 0) for g, v in (bucket or {}).items()}
        for guid in COMMERCIAL_GUIDS:
            out[month].setdefault(guid, 0.0)
    return out


def _save_odata_cache(year: int, ref_month: int, by_month: dict[int, dict[str, float]]) -> None:
    path = _odata_cache_path(year, ref_month)
    payload = {
        "cache_date": date.today().isoformat(),
        "cache_version": CACHE_VERSION,
        "source": "odata",
        "year": year,
        "ref_month": ref_month,
        "by_month": {str(m): by_month[m] for m in sorted(by_month)},
    }
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def find_valovaya_report(year: int, month: int) -> Path | None:
    """Найти выгрузку «валовая …» в temp по названию месяца, без фиксированного эталона."""
    folder = _temp_dir()
    if not folder.is_dir():
        return None
    month_name = MONTH_RU.get(month, "")
    candidates: list[Path] = []
    for path in folder.iterdir():
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
            continue
        name = path.name.lower()
        if "валовая" not in name:
            continue
        if "нет в план" in name:
            continue
        if month_name and month_name not in name:
            continue
        candidates.append(path)
    if not candidates:
        return None
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0]


def _row_values(ws, row: int) -> list[Any]:
    return [ws.cell(row, col).value for col in range(1, (ws.max_column or 1) + 1)]


def _find_vyruchka_col(values: list[Any]) -> int | None:
    found = None
    for idx, raw in enumerate(values, start=1):
        label = _norm_label(raw).lower()
        if not label:
            continue
        if label == "выручка" or (label.startswith("выручка") and "ндс" not in label):
            found = idx
            break
    return found


def _is_parent_or_header(label: str) -> bool:
    if not label:
        return True
    if label.lower() in {"подразделение", "итого"}:
        return True
    return bool(_PARENT_DEPT_RE.search(label))


def parse_valovaya_vyruchka_excel(path: Path) -> dict[str, Any] | None:
    """Прочитать «По подразделениям» / Выручка / Итого по подписям ячеек."""
    try:
        import openpyxl
    except ImportError:
        logger.warning("valovaya_vyruchka_otgruzki: openpyxl не установлен")
        return None
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
    except Exception:
        logger.exception("valovaya_vyruchka_otgruzki: не открыть %s", path)
        return None

    period_year = period_month = None
    section_row = None
    header_row = None
    vyruchka_col = None
    try:
        max_row = ws.max_row or 0
        for rno in range(1, max_row + 1):
            values = _row_values(ws, rno)
            texts = [_norm_label(v) for v in values if v is not None]
            joined = " ".join(texts)
            if period_year is None:
                match = _PERIOD_RE.search(joined)
                if match:
                    period_month = int(match.group(2))
                    period_year = int(match.group(3))
            if section_row is None and any("по подразделениям" in t.lower() for t in texts):
                section_row = rno
                continue
            if section_row is not None and header_row is None:
                col = _find_vyruchka_col(values)
                if col:
                    header_row = rno
                    vyruchka_col = col
                    break
        if not section_row or not header_row or not vyruchka_col:
            logger.warning("valovaya_vyruchka_otgruzki: нет блока По подразделениям в %s", path.name)
            return None

        by_guid: dict[str, float] = {}
        total = None
        for rno in range(header_row + 1, max_row + 1):
            label = _norm_label(ws.cell(rno, 1).value)
            if not label:
                if by_guid and total is not None:
                    break
                continue
            low = label.lower()
            if low.startswith("по менеджер"):
                break
            amount = _as_amount(ws.cell(rno, vyruchka_col).value)
            if amount is None:
                continue
            if low == "итого":
                total = amount
                break
            if _is_parent_or_header(label):
                continue
            guid = name_to_odata(label)
            if not guid:
                continue
            guid = normalize_commercial_dept_guid(guid)
            by_guid[guid] = by_guid.get(guid, 0.0) + float(amount)
    finally:
        wb.close()

    if total is None and by_guid:
        total = round(sum(by_guid.values()), 2)
    if total is None:
        return None
    return {
        "source_file": path.name,
        "year": period_year,
        "month": period_month,
        "total": round(float(total), 2),
        "by_guid": {g: round(v, 2) for g, v in by_guid.items()},
    }


def _excel_fact_for_month(year: int, month: int) -> dict[str, float] | None:
    path = find_valovaya_report(year, month)
    if path is None:
        return None
    parsed = parse_valovaya_vyruchka_excel(path)
    if not parsed:
        return None
    if parsed.get("year") and parsed.get("month"):
        if int(parsed["year"]) != year or int(parsed["month"]) != month:
            logger.info(
                "valovaya_vyruchka_otgruzki: %s период %s-%02d, нужен %s-%02d — пропуск",
                path.name, parsed.get("year"), parsed.get("month"), year, month,
            )
            return None
    by_guid = {str(g).lower(): float(v) for g, v in (parsed.get("by_guid") or {}).items()}
    out = {guid: float(by_guid.get(guid) or 0) for guid in COMMERCIAL_GUIDS}
    logger.info(
        "valovaya_vyruchka_otgruzki: %s итого %s по %s отделам",
        path.name, parsed.get("total"), sum(1 for v in out.values() if v),
    )
    return out


def get_vyruchka_otgruzki_fact_by_month(year: int, ref_month: int) -> dict[int, dict[str, float]]:
    """Помесячный факт из OData: {месяц: {guid отдела: сумма}}."""
    odata = _load_odata_cache(year, ref_month)
    if odata is None:
        odata = _fetch_odata_by_month(year, ref_month)
        if not any(sum(bucket.values()) for bucket in odata.values()):
            raise RuntimeError("OData: ВыручкаИСебестоимостьПродаж пуст или не ответил")
        _save_odata_cache(year, ref_month, odata)
    result: dict[int, dict[str, float]] = {}
    for month in range(1, int(ref_month) + 1):
        excel_map = _excel_fact_for_month(year, month)
        if excel_map is not None:
            result[month] = excel_map
            continue
        bucket = odata.get(month) or _empty_by_guid()
        result[month] = {guid: float(bucket.get(guid) or 0) for guid in COMMERCIAL_GUIDS}
        for guid, amt in bucket.items():
            result[month][str(guid).lower()] = float(amt or 0)
    return result


def named_fact_by_dept(by_guid: dict[str, float]) -> dict[str, float]:
    return {
        name: round(float(by_guid.get(guid) or 0), 2)
        for guid, name in DEPT_GUID_TO_DZ_NAME.items()
    }
